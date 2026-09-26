"""Excel/CSV export and import helpers."""

import csv
import io
from collections.abc import Iterable, Sequence
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..pipeline import stage_label

HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(bold=True, color="FFFFFF")

# (header, attribute) pairs for the candidate tracker sheet.
CANDIDATE_COLUMNS: tuple[tuple[str, str], ...] = (
    ("Candidate Name", "full_name"),
    ("Phone", "phone"),
    ("Email", "email"),
    ("Current Company", "current_company"),
    ("Current Designation", "current_designation"),
    ("Total Exp (yrs)", "total_experience"),
    ("Relevant Exp (yrs)", "relevant_experience"),
    ("Current CTC (LPA)", "current_ctc"),
    ("Expected CTC (LPA)", "expected_ctc"),
    ("Offered CTC (LPA)", "offered_ctc"),
    ("Notice Period (days)", "notice_period_days"),
    ("Current Location", "current_location"),
    ("Preferred Location", "preferred_location"),
    ("Skills", "skills"),
    ("Source", "source"),
    ("Resume Link", "resume_url"),
    ("Stage", "stage"),
    ("Joining Date", "joining_date"),
    ("Remarks", "remarks"),
)

_ALIASES = {
    "name": "full_name",
    "candidate": "full_name",
    "candidate name": "full_name",
    "mobile": "phone",
    "contact": "phone",
    "phone number": "phone",
    "mail": "email",
    "email id": "email",
    "company": "current_company",
    "designation": "current_designation",
    "experience": "total_experience",
    "exp": "total_experience",
    "total exp": "total_experience",
    "ctc": "current_ctc",
    "cctc": "current_ctc",
    "ectc": "expected_ctc",
    "notice": "notice_period_days",
    "notice period": "notice_period_days",
    "location": "current_location",
    "resume": "resume_url",
    "status": "stage",
    "comments": "remarks",
    "notes": "remarks",
}


def _normalise(header: str) -> str:
    return " ".join(str(header).strip().lower().replace("_", " ").split())


def header_to_field(header: Any) -> str | None:
    if header is None:
        return None
    key = _normalise(header)
    for label, attr in CANDIDATE_COLUMNS:
        if key in (_normalise(label), _normalise(attr)):
            return attr
    stripped = key.split("(")[0].strip()
    for label, attr in CANDIDATE_COLUMNS:
        if stripped == _normalise(label).split("(")[0].strip():
            return attr
    return _ALIASES.get(key) or _ALIASES.get(stripped)


def build_workbook(sheets: Sequence[tuple[str, Sequence[str], Iterable[Sequence[Any]]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title=title[:31])
        ws.append(list(headers))
        for cell in ws[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center")
        widths = [len(str(h)) for h in headers]
        for row in rows:
            values = list(row)
            ws.append(values)
            for i, v in enumerate(values):
                widths[i] = max(widths[i], min(len(str(v)) if v is not None else 0, 60))
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width + 3
        ws.freeze_panes = "B2"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def candidate_rows(candidates: Iterable[Any]) -> list[list[Any]]:
    rows = []
    for c in candidates:
        row = []
        for _, attr in CANDIDATE_COLUMNS:
            value = getattr(c, attr)
            if attr == "stage":
                value = stage_label(value)
            row.append(value)
        row.append(c.recruiter.full_name if c.recruiter else None)
        row.append(c.created_at.date() if c.created_at else None)
        rows.append(row)
    return rows


CANDIDATE_HEADERS = [label for label, _ in CANDIDATE_COLUMNS] + ["Recruiter", "Added On"]


def read_table(filename: str, content: bytes) -> list[dict[str, Any]]:
    """Parse an uploaded .xlsx or .csv into a list of {field: value} dicts."""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = list(wb.active.iter_rows(values_only=True))
    else:
        text = content.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        return []
    fields = [header_to_field(h) for h in rows[0]]
    records = []
    for raw in rows[1:]:
        if raw is None or all(v in (None, "") for v in raw):
            continue
        record = {}
        for field, value in zip(fields, raw):
            if field and value not in (None, ""):
                record[field] = value.strip() if isinstance(value, str) else value
        records.append(record)
    return records
