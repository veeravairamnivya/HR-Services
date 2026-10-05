import io

from openpyxl import load_workbook


def _add(client, headers, position_id, **fields):
    body = {"full_name": "Asha Kumar", "email": "asha@example.com", "phone": "9876543210", **fields}
    return client.post(f"/api/positions/{position_id}/candidates", headers=headers, json=body)


def test_sheet_add_edit_and_duplicates(client, setup):
    pid = setup["position"]["id"]
    r = _add(client, setup["rita"], pid, current_ctc=10, expected_ctc="")
    assert r.status_code == 201, r.text
    cand = r.json()
    assert cand["stage"] == "sourced" and cand["recruiter"]["full_name"] == "Rita Recruiter"

    dup = _add(client, setup["rita"], pid, full_name="Someone Else", email="other@example.com", phone="+91 98765 43210")
    assert dup.status_code == 409

    r = client.patch(f"/api/candidates/{cand['id']}", headers=setup["rita"], json={"expected_ctc": 14.5, "remarks": "Good fit"})
    assert r.status_code == 200 and r.json()["expected_ctc"] == 14.5

    sheet = client.get(f"/api/positions/{pid}/candidates", headers=setup["rita"]).json()
    assert [c["full_name"] for c in sheet] == ["Asha Kumar"]
    matches = client.get("/api/candidates/duplicates", headers=setup["rita"], params={"phone": "98765-43210"}).json()
    assert len(matches) == 1


def test_unassigned_recruiter_cannot_edit(client, setup):
    pid = setup["position"]["id"]
    cand = _add(client, setup["rita"], pid).json()
    r = client.patch(f"/api/candidates/{cand['id']}", headers=setup["omar"], json={"remarks": "hijack"})
    assert r.status_code == 403
    # ...but can add their own candidate and edit it.
    own = _add(client, setup["omar"], pid, full_name="Omar Cand", email="oc@example.com", phone="9000000001").json()
    assert client.patch(f"/api/candidates/{own['id']}", headers=setup["omar"], json={"remarks": "ok"}).status_code == 200


def test_interview_rounds_drive_stage(client, setup):
    pid = setup["position"]["id"]
    h = setup["rita"]
    cand = _add(client, h, pid, stage="shortlisted").json()

    r1 = client.post(f"/api/candidates/{cand['id']}/interviews", headers=h,
                     json={"scheduled_at": "2030-01-10T10:00:00Z", "interviewer": "Panel"})
    assert r1.status_code == 201 and r1.json()["round_number"] == 1
    detail = client.get(f"/api/candidates/{cand['id']}", headers=h).json()
    assert detail["stage"] == "interview_scheduled"
    assert detail["next_interview"]["id"] == r1.json()["id"]

    client.patch(f"/api/interviews/{r1.json()['id']}", headers=h, json={"result": "selected", "feedback": "Strong"})
    assert client.get(f"/api/candidates/{cand['id']}", headers=h).json()["stage"] == "round1_selected"

    r2 = client.post(f"/api/candidates/{cand['id']}/interviews", headers=h, json={"scheduled_at": "2030-01-12T10:00:00Z"})
    assert r2.json()["round_number"] == 2
    # scheduling round 2 must not move the candidate backwards
    assert client.get(f"/api/candidates/{cand['id']}", headers=h).json()["stage"] == "round1_selected"
    client.patch(f"/api/interviews/{r2.json()['id']}", headers=h, json={"result": "selected"})
    assert client.get(f"/api/candidates/{cand['id']}", headers=h).json()["stage"] == "round2_selected"

    for stage in ("hr_discussion", "offer_released"):
        r = client.post(f"/api/candidates/{cand['id']}/stage", headers=h, json={"stage": stage})
        assert r.status_code == 200
    detail = client.get(f"/api/candidates/{cand['id']}", headers=h).json()
    assert [x["to_stage"] for x in detail["history"]][:3] == ["offer_released", "hr_discussion", "round2_selected"]


def test_rejection_via_interview(client, setup):
    h = setup["rita"]
    cand = _add(client, h, setup["position"]["id"]).json()
    iv = client.post(f"/api/candidates/{cand['id']}/interviews", headers=h, json={}).json()
    client.patch(f"/api/interviews/{iv['id']}", headers=h, json={"result": "rejected"})
    assert client.get(f"/api/candidates/{cand['id']}", headers=h).json()["stage"] == "rejected"


def test_bulk_stage_and_dashboard(client, admin, setup):
    h = setup["rita"]
    pid = setup["position"]["id"]
    ids = [
        _add(client, h, pid, full_name=f"Cand {i}", email=f"c{i}@example.com", phone=f"90000000{i:02d}").json()["id"]
        for i in range(3)
    ]
    r = client.post("/api/candidates/bulk-stage", headers=h, json={"candidate_ids": ids, "stage": "shortlisted"})
    assert r.json() == {"updated": 3, "denied": 0}

    dash = client.get("/api/dashboard", headers=admin).json()
    assert dash["kpis"]["total_candidates"] == 3
    assert dash["kpis"]["open_positions"] == 1
    funnel = {f["stage"]: f["count"] for f in dash["funnel"]}
    assert funnel["shortlisted"] == 3 and funnel["interview_scheduled"] == 0

    perf = client.get("/api/reports/recruiters", headers=admin).json()["rows"]
    rita = next(r for r in perf if r["name"] == "Rita Recruiter")
    assert rita["added"] == 3 and rita["shortlisted"] == 3

    positions = client.get("/api/positions", headers=admin).json()
    assert positions[0]["candidate_count"] == 3 and positions[0]["stage_counts"] == {"shortlisted": 3}
    clients = client.get("/api/clients", headers=admin).json()
    assert clients[0]["open_positions"] == 1 and clients[0]["total_candidates"] == 3


def test_export_and_import(client, setup):
    h = setup["rita"]
    pid = setup["position"]["id"]
    csv_content = (
        "Name,Mobile,Email,Company,Experience,Status,Source\n"
        "Ravi Teja,9111111111,ravi@example.com,TCS,4,Shortlisted,LinkedIn\n"
        "Bad Row,9222222222,not-an-email,,,,\n"
        "Ravi Again,9111111111,,,,,\n"
    )
    r = client.post(f"/api/positions/{pid}/candidates/import", headers=h,
                    files={"file": ("cands.csv", csv_content, "text/csv")})
    assert r.status_code == 200, r.text
    assert r.json()["created"] == 1 and r.json()["skipped"] == 2
    sheet = client.get(f"/api/positions/{pid}/candidates", headers=h).json()
    assert sheet[0]["stage"] == "shortlisted" and sheet[0]["source"] == "linkedin"

    r = client.get(f"/api/positions/{pid}/candidates/export", headers=h)
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws["A1"].value == "Candidate Name" and ws["A2"].value == "Ravi Teja"

    for kind in ("recruiters", "clients", "positions", "attendance", "candidates"):
        assert client.get(f"/api/reports/export/{kind}", headers=h).status_code == 200
    assert client.get("/api/reports/export/daily", headers=h).status_code == 403


def test_closed_position_rejects_candidates(client, admin, setup):
    pid = setup["position"]["id"]
    client.patch(f"/api/positions/{pid}", headers=admin, json={"status": "closed"})
    assert _add(client, setup["rita"], pid).status_code == 400


def test_links_must_be_http(client, setup):
    h = setup["rita"]
    pid = setup["position"]["id"]
    assert _add(client, h, pid, resume_url="javascript:alert(1)").status_code == 422
    r = _add(client, h, pid, resume_url="drive.google.com/file/abc")
    assert r.status_code == 201 and r.json()["resume_url"] == "https://drive.google.com/file/abc"


def test_dashboard_cards_match_their_candidate_lists(client, admin, setup):
    """Every dashboard number links to /candidates with filters that must return exactly that many rows."""
    h = setup["rita"]
    pid = setup["position"]["id"]
    stages = ["sourced", "shortlisted", "offer_released", "joined", "rejected"]
    for i, stage in enumerate(stages):
        r = _add(client, h, pid, full_name=f"Card {i}", email=f"card{i}@example.com", phone=f"91111111{i:02d}")
        assert r.status_code == 201
        client.post(f"/api/candidates/{r.json()['id']}/stage", headers=h, json={"stage": stage})

    dash = client.get("/api/dashboard", headers=admin).json()
    today = dash["today"]
    month = today[:8] + "01"
    active = "sourced,screening,shortlisted,interview_scheduled,round1_selected,round2_selected,round3_selected," \
             "hr_discussion,offer_released,offer_accepted,on_hold"

    def total(**params):
        r = client.get("/api/candidates", headers=admin, params=params)
        assert r.status_code == 200, r.text
        return r.json()["total"]

    k = dash["kpis"]
    assert total(stage=active) == k["active_pipeline"] == 3
    assert total(added_from=month, added_to=today) == k["added_month"] == 5
    assert total(reached="offer_released", reached_from=month, reached_to=today) == k["offers_month"] == 1
    assert total(reached="joined", reached_from=month, reached_to=today) == k["joined_month"] == 1
    for tile in dash["stage_distribution"]:
        assert total(stage=tile["stage"], position_status="active") == tile["count"]
